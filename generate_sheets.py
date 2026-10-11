import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()

# ==========================================
# SHEET 1: AGRISHARE EQUIPMENT DATABASE
# Columns: ID | Equipment | Owner | Village | Rate/hr | Status | Best Month | Available Date | Contact Number
# ==========================================
ws1 = wb.active
ws1.title = 'AGRISHARE_DATABASE'
ws1.views.sheetView[0].showGridLines = True

# Title
ws1.merge_cells('A1:I1')
title_cell = ws1['A1']
title_cell.value = 'AGRISHARE EQUIPMENT DATABASE — Master Data Registry'
title_cell.font = Font(name='Arial', size=14, bold=True, color='FFFFFF')
title_cell.fill = PatternFill(start_color='237A3B', end_color='237A3B', fill_type='solid')
title_cell.alignment = Alignment(horizontal='center', vertical='center')
ws1.row_dimensions[1].height = 34

# Subtitle / Account Info
ws1.merge_cells('A2:I2')
sub_cell = ws1['A2']
sub_cell.value = 'A Smart Agricultural Equipment Sharing Platform | Viksit Bharat Rural Mechanization'
sub_cell.font = Font(name='Arial', size=10, italic=True, color='FFFFFF')
sub_cell.fill = PatternFill(start_color='164E24', end_color='164E24', fill_type='solid')
sub_cell.alignment = Alignment(horizontal='center', vertical='center')
ws1.row_dimensions[2].height = 22

# Blank spacer
ws1.append([])
ws1.row_dimensions[3].height = 6

# Exact 9 Column Headers matching Master Prompt Specification:
headers = [
    'ID',
    'Equipment',
    'Owner',
    'Village',
    'Rate/hr',
    'Status',
    'Best Month',
    'Available Date',
    'Contact Number'
]
ws1.append(headers)
ws1.row_dimensions[4].height = 28

header_fill = PatternFill(start_color='237A3B', end_color='237A3B', fill_type='solid')
header_font = Font(name='Arial', size=11, bold=True, color='FFFFFF')
thin_border = Border(
    left=Side(style='thin', color='D1D5DB'),
    right=Side(style='thin', color='D1D5DB'),
    top=Side(style='thin', color='D1D5DB'),
    bottom=Side(style='thin', color='D1D5DB')
)

for col_num in range(1, len(headers) + 1):
    c = ws1.cell(row=4, column=col_num)
    c.fill = header_fill
    c.font = header_font
    c.alignment = Alignment(horizontal='center', vertical='center')
    c.border = thin_border

# Exact 14 sample records matching Section 6
data = [
    ('A01', 'Rotavator', 'Farmer A', 'Village A', 800, 'Available', 'April–May', '2026-10-15', '9820112233'),
    ('A02', 'Tractor', 'Farmer B', 'Village B', 1000, 'Available', 'April–June', '2026-10-16', '9820334455'),
    ('A03', 'Seed Drill', 'Farmer C', 'Village A', 600, 'Booked', 'June–July', '2026-10-17', '9820445566'),
    ('A04', 'Sprayer', 'Farmer D', 'Village C', 250, 'Available', 'July–September', '2026-10-18', '9820123456'),
    ('A05', 'Harvester', 'Farmer E', 'Village B', 1500, 'Available', 'October–November', '2026-10-22', '9820667788'),
    ('A06', 'Thresher', 'Farmer F', 'Village A', 700, 'Available', 'October–December', '2026-10-25', '9820551122'),
    ('A07', 'Rotavator', 'Farmer G', 'Village C', 750, 'Available', 'November–December', '2026-10-20', '9820773344'),
    ('A08', 'Tractor', 'Farmer H', 'Village A', 1100, 'Available', 'October–November', '2026-10-21', '9820884455'),
    ('A09', 'Seed Drill', 'Farmer I', 'Village B', 650, 'Available', 'November–December', '2026-10-23', '9820995566'),
    ('A10', 'Sprayer', 'Farmer J', 'Village B', 280, 'Booked', 'January–February', '2026-10-19', '9820226677'),
    ('A11', 'Harvester', 'Farmer K', 'Village A', 1600, 'Available', 'March–April', '2026-10-24', '9820337788'),
    ('A12', 'Thresher', 'Farmer L', 'Village C', 750, 'Available', 'March–April', '2026-10-26', '9820448899'),
    ('A13', 'Laser Land Leveller', 'Farmer M', 'Village A', 950, 'Available', 'April–May', '2026-10-27', '9820556677'),
    ('A14', 'Power Tiller', 'Farmer N', 'Village B', 500, 'Available', 'June–July', '2026-10-28', '9820667788'),
]

available_fill = PatternFill(start_color='DCFCE7', end_color='DCFCE7', fill_type='solid')
available_font = Font(name='Arial', size=10, bold=True, color='15803D')
booked_fill = PatternFill(start_color='FEE2E2', end_color='FEE2E2', fill_type='solid')
booked_font = Font(name='Arial', size=10, bold=True, color='B91C1C')
regular_font = Font(name='Arial', size=10)

for row_idx, row_data in enumerate(data, start=5):
    ws1.row_dimensions[row_idx].height = 24
    for col_idx, val in enumerate(row_data, start=1):
        c = ws1.cell(row=row_idx, column=col_idx, value=val)
        c.font = regular_font
        c.border = thin_border
        
        # ID column center
        if col_idx == 1:
            c.alignment = Alignment(horizontal='center', vertical='center')
            c.font = Font(name='Arial', size=10, bold=True)
        # Rate column currency format
        elif col_idx == 5:
            c.alignment = Alignment(horizontal='right', vertical='center')
            c.number_format = '₹#,##0'
        # Status column styling
        elif col_idx == 6:
            c.alignment = Alignment(horizontal='center', vertical='center')
            if val == 'Available':
                c.fill = available_fill
                c.font = available_font
            else:
                c.fill = booked_fill
                c.font = booked_font
        # Available Date
        elif col_idx == 8:
            c.alignment = Alignment(horizontal='center', vertical='center')
        else:
            c.alignment = Alignment(horizontal='left', vertical='center')

# Auto-adjust column widths
for col in ws1.columns:
    max_len = 0
    col_letter = get_column_letter(col[0].column)
    for cell in col:
        if cell.row > 2 and cell.value:
            max_len = max(max_len, len(str(cell.value)))
    ws1.column_dimensions[col_letter].width = max(max_len + 4, 12)

# Save Workbook
output_path = 'AGRISHARE_EQUIPMENT_DATABASE.xlsx'
wb.save(output_path)
print(f"Successfully generated {output_path} with 14 simulated records.")
